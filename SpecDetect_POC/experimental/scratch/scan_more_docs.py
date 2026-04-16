import os
from pypdf import PdfReader
from docx import Document

files = [
    r'D:\arvin\YL_workapace\超短波监测管理一体化服务接口规范\soap协议\GWJ001-2015超短波监测管理一体化服务接口规范 平台架构部分.pdf',
    r'D:\arvin\YL_workapace\超短波监测管理一体化服务接口规范\soap协议\GWJ002-2015超短波监测管理一体化服务接口规范 服务和接口部分.pdf',
    r'D:\arvin\YL_workapace\超短波监测管理一体化服务接口规范\soap协议\GWJ003-2015超短波监测管理一体化服务接口规范 设备操作服务部分.pdf',
    r'D:\arvin\YL_workapace\超短波监测管理一体化服务接口规范\soap协议\附件1：《超短波监测管理一体化平台技术规范 第3部分：设备操作服务》SOAP报文结构补充说明1.docx',
    r'D:\arvin\YL_workapace\超短波监测管理一体化服务接口规范\soap协议\工信部无2016379号-1超短波监测管理服务接口规范.docx',
]

keywords = ['压缩', '量化', '归一', '滤波', '切片', '分片', '转发', 'Atom', 'atom',
            '变换', '采样', '抽取', '降采样', '门限', '平均', '平滑', '预处理',
            'streamsrc', 'RMCPTP', 'STC', '0xEEEEEEEE', '0xEFFF', 'LEADER',
            '实时', '推送', '订阅', 'subscribe', '数据流', '流式', '频谱数据',
            '频率序号', '频率数量', 'nOffset', '检波', 'squelch', '静噪']

out = open('more_docs_hits.txt', 'w', encoding='utf-8')
def w(s): out.write(s + '\n')

for path in files:
    name = os.path.basename(path)
    w('=' * 70)
    w('FILE: ' + name)
    w('=' * 70)
    try:
        if path.lower().endswith('.pdf'):
            r = PdfReader(path)
            for pi, pg in enumerate(r.pages):
                txt = (pg.extract_text() or '').replace('\x00', '')
                for ln in txt.split('\n'):
                    s = ln.strip()
                    if s and any(k in s for k in keywords):
                        w(f'[p{pi+1}] {s}')
        else:
            d = Document(path)
            for i, para in enumerate(d.paragraphs):
                t = para.text.strip()
                if t and any(k in t for k in keywords):
                    w(f'[p{i}] {t}')
            for ti, tbl in enumerate(d.tables):
                for ri, row in enumerate(tbl.rows):
                    joined = ' | '.join(c.text.strip() for c in row.cells)
                    if any(k in joined for k in keywords):
                        w(f'[t{ti}r{ri}] {joined}')
    except Exception as e:
        w(f'ERROR: {e}')
out.close()
print('done')
