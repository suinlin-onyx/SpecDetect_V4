import sys
from docx import Document

paths = [
    r'D:\arvin\YL_workapace\超短波监测管理一体化服务接口规范\soap协议\GWJ004-2015超短波监测管理一体化服务接口规范 数据服务部分.docx',
    r'D:\arvin\YL_workapace\超短波监测管理一体化服务接口规范\soap协议\GWJ006-2016超短波频段监测基础数据存储结构技术规范.docx',
]

keywords = ['压缩', '量化', '归一', '滤波', '切片', '分片', '转发', 'atom', 'Atom',
            '变换', '采样', '抽取', '降采样', 'dBuV', 'dBm', '电平', '无效',
            '门限', '平均', '平滑', '预处理', '处理', '编码']

out = open('docx_keywords.txt', 'w', encoding='utf-8')
def p(s):
    out.write(s + '\n')
for path in paths:
    name = path.split('\\')[-1]
    p('=' * 70)
    p('FILE: ' + name)
    p('=' * 70)
    d = Document(path)
    for i, para in enumerate(d.paragraphs):
        t = para.text.strip()
        if not t:
            continue
        if any(k in t for k in keywords):
            p(f'[p{i}] {t}')
    # tables
    for ti, tbl in enumerate(d.tables):
        for ri, row in enumerate(tbl.rows):
            cells = [c.text.strip() for c in row.cells]
            joined = ' | '.join(cells)
            if any(k in joined for k in keywords):
                p(f'[t{ti}r{ri}] {joined}')
